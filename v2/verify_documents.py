import pathlib,json,re,hashlib,math,subprocess
from docx import Document
import fitz
from PIL import Image,ImageDraw
P=pathlib.Path(__file__).resolve().parents[1];D=P/'dist';Q=P/'qa-v2';Q.mkdir(exist_ok=True);C=json.loads((D/'course.json').read_text());errors=[];stats=[]
norm=lambda s:re.sub(r'\s+',' ',re.sub(r'(?<=\w)([-–])\s+(?=\w)',r'\1',s)).strip()
for name in ['Curriculum.pdf','Questions.pdf','Practice.docx']:
 a=P.parent/'upload'/('Unit1_'+name);b=D/'units/Unit1'/('Unit1_'+name)
 assert hashlib.sha256(a.read_bytes()).digest()==hashlib.sha256(b.read_bytes()).digest()
for l in C['lessons'][1:]:
 n=l['number'];row={'unit':n};assert len(l['vocab'])==20 and len(l['questions'])==19
 assert sum(q['points'] for q in l['questions'])==100
 word=Document(D/l['downloads'][2]);raw=' '.join(p.text for p in word.paragraphs)+' '+' '.join(c.text for t in word.tables for r in t.rows for c in r.cells)
 expectedcodes=[q['code'] for q in l['questions']]
 for q in l['questions']:
  assert norm(q['prompt']) in norm(raw),(n,q['code'],'Word prompt mismatch')
  if q['type']=='mc':
   for o in q['options']:assert norm(o) in norm(raw)
   # Locate the 2x2 choice table and prove the row order is A/B, then C/D.
   tables=[t for t in word.tables if len(t.rows)==2 and len(t.columns)==2 and norm(t.cell(0,0).text)==norm('A. '+q['options'][0])]
   assert len(tables)==1,(n,q['code'],'choice table')
   assert [norm(c.text) for r in tables[0].rows for c in r.cells]==[norm(chr(65+i)+'. '+o) for i,o in enumerate(q['options'])]
 render=Q/'docx'/f'Unit{n}'/f'Unit{n}_Practice.pdf';assert render.exists(),n
 files=[D/l['downloads'][0],D/l['downloads'][1],render];thumbs=[]
 for j,f in enumerate(files):
  doc=fitz.open(f);row[['curriculumPages','questionPages','wordPages'][j]]=len(doc)
  text='\n'.join(p.get_text() for p in doc)
  if j in [1,2]:
   for q in l['questions']:
    assert norm(q['prompt']) in norm(text),(n,j,q['code'],'render prompt mismatch')
    if q['type']=='mc':
     pages=[p for p in doc if norm(q['prompt']) in norm(p.get_text())];assert len(pages)==1,(n,j,q['code'],'stem split')
     page=pages[0]
     for k,o in enumerate(q['options']):assert norm(chr(65+k)+'. '+o) in norm(page.get_text()),(n,j,q['code'],'choices split')
   assert norm(l['reading']) in norm(text),(n,j,'missing passage')
  for i,page in enumerate(doc):
   for b in page.get_text('dict')['blocks']:
    for line in b.get('lines',[]):
     for s in line['spans']:
      x0,y0,x1,y1=s['bbox']
      if x0<35 or x1>page.rect.width-30 or y0<16 or y1>page.rect.height-15:errors.append((n,j,i+1,s['text'],s['bbox']))
      if s['size']<8:errors.append((n,j,i+1,'font too small',s['size']))
      if '\ufffd' in s['text']:errors.append((n,j,i+1,'replacement glyph'))
   if j==2:
    body=[line for line in page.get_text().splitlines() if not line.startswith('Catholic English | Unit') and not line.isdigit()]
    if not any(x.strip() for x in body):errors.append((n,j,i+1,'blank page'))
   pix=page.get_pixmap(matrix=fitz.Matrix(.29,.29),alpha=False);im=Image.frombytes('RGB',(pix.width,pix.height),pix.samples);thumbs.append((j,i+1,im))
  if n in [2,12,20,26] and j in [0,1]:
   for i in [0,min(2,len(doc)-1)]:doc[i].get_pixmap(matrix=fitz.Matrix(1.2,1.2)).save(str(Q/f'Unit{n}_{j}_p{i+1}.png'))
 # One contact sheet per unit contains every page of all three documents.
 cols=5;tw,th=190,275;rows=math.ceil(len(thumbs)/cols);sheet=Image.new('RGB',(cols*tw,rows*th),'#dddddd');draw=ImageDraw.Draw(sheet)
 for k,(j,page,im) in enumerate(thumbs):
  x=(k%cols)*tw;y=(k//cols)*th;sheet.paste(im,(x+8,y+19));draw.text((x+8,y+4),['Curriculum','Questions','Word'][j]+' '+str(page),fill='black')
 sheet.save(Q/f'Unit{n:02}_all_pages.jpg',quality=85);stats.append(row)
report={'newUnits':25,'newDocuments':75,'pdfs':50,'docx':25,'unit1FilesUnchanged':True,'wordRenderComplete':True,'allPromptsAndReadingsMatch':True,'allMCOptionsTwoByTwoAndWithStem':True,'errors':errors,'pages':stats,'browserQA':'not available'}
(Q/'report-v2.json').write_text(json.dumps(report,ensure_ascii=False,indent=2));print(json.dumps({k:v for k,v in report.items() if k!='pages'},ensure_ascii=False));print('Pages total',sum(sum(x[k] for k in ['curriculumPages','questionPages','wordPages']) for x in stats))
assert not errors
