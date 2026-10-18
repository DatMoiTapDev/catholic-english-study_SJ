import pathlib,json,html,re
from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Table,TableStyle,PageBreak,KeepTogether
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
P=pathlib.Path(__file__).resolve().parents[1];D=P/'dist';C=json.loads((D/'course.json').read_text());W=A4[0]-92
pdfmetrics.registerFont(TTFont('Book','/usr/share/fonts/truetype/dejavu/DejaVuSerif.ttf'));pdfmetrics.registerFont(TTFont('BookB','/usr/share/fonts/truetype/dejavu/DejaVuSerif-Bold.ttf'));pdfmetrics.registerFont(TTFont('IPA','/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'))
pdfmetrics.registerFontFamily('Book',normal='Book',bold='BookB',italic='Book',boldItalic='BookB')
S={
'b':ParagraphStyle('b',fontName='Book',fontSize=11,leading=16,spaceAfter=8),
's':ParagraphStyle('s',fontName='Book',fontSize=9,leading=13,spaceAfter=6),
'cell':ParagraphStyle('cell',fontName='Book',fontSize=10,leading=14,spaceAfter=4),
'ipa':ParagraphStyle('ipa',fontName='IPA',fontSize=10,leading=14,spaceAfter=4),
'h':ParagraphStyle('h',fontName='BookB',fontSize=15,leading=21,spaceBefore=14,spaceAfter=11,keepWithNext=True),
't':ParagraphStyle('t',fontName='BookB',fontSize=19,leading=25,spaceAfter=12),
'v':ParagraphStyle('v',fontName='BookB',fontSize=11,leading=15,spaceAfter=5,keepWithNext=True)}
def para(t,style='b'):return Paragraph(html.escape(str(t)).replace('\n','<br/>'),S[style])
def head(t):return para(t,'h')
def table(rows,widths,header=True,grid=True,styles=None):
 data=[]
 for i,row in enumerate(rows):
  data.append([x if isinstance(x,(list,Paragraph)) else para(x,'cell') for x in row])
 t=Table(data,colWidths=widths,repeatRows=1 if header else 0,hAlign='LEFT')
 ops=[('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),7),('RIGHTPADDING',(0,0),(-1,-1),7),('TOPPADDING',(0,0),(-1,-1),6),('BOTTOMPADDING',(0,0),(-1,-1),5)]
 if grid:ops +=[('GRID',(0,0),(-1,-1),.35,colors.HexColor('#bbbbbb'))]
 if header:ops +=[('BACKGROUND',(0,0),(-1,0),colors.HexColor('#eeeeee'))]
 t.setStyle(TableStyle(ops+(styles or [])));return t

def bilingual(pairs):return table([[para(a),para(b)] for a,b in pairs],[W/2,W/2],False,False)
def footer(cv,doc):
 cv.setFont('Book',8);cv.drawString(46,25,f'Catholic English | Unit {doc.uid} | B1 core and optional extension');cv.drawRightString(A4[0]-46,25,str(doc.page))
def writepdf(path,story,uid):
 doc=SimpleDocTemplate(str(path),pagesize=A4,leftMargin=46,rightMargin=46,topMargin=42,bottomMargin=44,title=f'Catholic English Unit {uid}',author='Học liệu bổ trợ theo mẫu Unit 1');doc.uid=uid;doc.build(story,onFirstPage=footer,onLaterPages=footer)
GROUPS=[('A','Reading','Reading comprehension','R01–R05'),('B','Terminology','Catholic terminology','V01–V04'),('C','Grammar','Grammar in context','G01–G03'),('D','Speaking','Speaking and community communication','S01–S02'),('E','Writing','Academic and practical writing','W01'),('F','Critical reading','Critical reading','C01'),('G','Listening','Listening in context','L01–L03')]
def marking(l):return [['TASK / NHÓM BÀI','MÃ CÂU','ĐIỂM']]+[[a+'. '+name,codes,str(sum(q['points'] for q in l['questions'] if q['skill']==skill))] for a,skill,name,codes in GROUPS]+[['TOTAL','19 nhiệm vụ','100']]
def optiontable(q):return table([[chr(65+k)+'. '+q['options'][k] for k in [0,1]],[chr(65+k)+'. '+q['options'][k] for k in [2,3]]],[W/2,W/2],False)
def qpdf(q):
 b=[para(f"{q['code']} | {q['prompt']} ({q['points']} points)")]
 if q['type']=='mc':b+=[optiontable(q),Spacer(1,12)]
 else:b.append(Spacer(1,9))
 return KeepTogether(b)

def source_items(l):
 ids=list(dict.fromkeys([l['source'],'girm6' if (l['number']-1)//2==5 else 'girm3' if (l['number']-1)//2==10 else 'examen' if (l['number']-1)//2==9 else 'ccc']))
 if l['source']=='sj':ids.append('magis')
 if l['source']=='lg' and l['number']>=23:ids.append('gs')
 return [(C['sources'][i]['title'],C['sources'][i]['url']) for i in ids]

def curriculum(l):
 n=l['number'];s=[para(f'UNIT {n:02} {l["title"].upper()}','t'),para('Bilingual Catholic English self study textbook','s'),para('B1 nền tảng • B2 mở rộng tùy năng lực • Phần '+l['part']+' • '+l['date'],'s'),para('BILINGUAL READING GUIDE  ENGLISH LEFT  TIẾNG VIỆT RIGHT','v'),table([['Course / Người học','Scope / Mục tiêu'],['Catholic Vocabulary and Ecclesial Language',l['focus']],[f'Lesson code CE-L{n:02}',l['selfStudyMinutes']],['Source priority',l['reference']]],[W*.46,W*.54]),head('1 Introduction and Learning Objectives'),bilingual([l['intro']]),table([['Learning objective','Observable outcome'],['Define and distinguish','Explain four conceptual distinctions using correct examples.'],['Read for argument','Identify main ideas, details, inference, and a word in context.'],['Speak and write','Produce a speaking script, six-turn dialogue, and a purposeful 150–180-word text.'],['Listen in context','Identify the purpose and practical details in a playable audio recording.']],[W*.34,W*.66]),para('Thiết kế bổ trợ theo syllabus, chưa được giảng viên phê duyệt. B2 chỉ chỉ phần luyện nâng cao, không chứng nhận trình độ. Bài làm dài có thể hoàn thành ngoài giờ học.','s'),PageBreak(),head('2 Essential Catholic Vocabulary'),para('MUST KNOW = nắm chắc • SHOULD KNOW = nên dùng được • ADVANCED = mở rộng. IPA ưu tiên giọng Anh Mỹ; cách đọc các từ Latinh là tham khảo theo lối đọc tiếng Anh.','s')]
 rows=[['Term / IPA / POS','Nghĩa tiếng Việt','Definition / Example / Collocation','Priority']]
 for v in l['vocab']:
  rows.append([[para(v['word'],'v'),para(v['ipa'],'ipa'),para(v['pos'],'s')],v['meaning'],[para(v['definition'],'cell'),para('Example: '+v['example'],'cell'),para('Collocation: '+v['collocation'],'s')],para(v['priority'],'s')])
 s +=[table(rows,[W*.21,W*.18,W*.47,W*.14]),PageBreak(),head('3 Key Concepts and Distinctions')]
 for pair,en,vi in l['distinctions']:
  s +=[head(pair.replace('/',' and ')),bilingual([(en,vi)]),para('Try it: explain the difference without translating word by word. Then choose a sentence from the reading that makes the context clear.','s')]
 s +=[bilingual([(l['core'][l['number']%2][0],l['core'][l['number']%2][1])]),PageBreak(),head('4 Grammar and Sentence Patterns'),table([['Structure','Example','Usage / Giải thích']]+[[a,b,c] for a,b,c in l['patterns']],[W*.26,W*.40,W*.34]),head('Guided language practice'),bilingual([('Start with a clear subject and verb. Add a reason or contrast only when the basic message is accurate. A correct short sentence is more useful than a long sentence with uncertain meaning.','Bắt đầu bằng chủ ngữ và động từ rõ. Chỉ thêm lý do hoặc tương phản khi ý cơ bản đã đúng. Câu ngắn đúng hữu ích hơn câu dài mà nghĩa chưa chắc.')]),para(l['grammar']),head('Pronunciation and self correction'),para(l['pronunciation']),para('Nghe lại bản ghi của mình: người nghe có nhận ra từ khóa không, câu có ngắt theo ý không? Không thể chấm phát âm chỉ từ kịch bản viết. IPA cũ trong phần ghi chú có thể theo giọng Anh; bảng từ là tham khảo Anh Mỹ.','s'),head('Common errors'),para('Nhầm các vế khái niệm trong phần 3; dùng từ đúng chính tả nhưng sai vai trò; chép cấu trúc mà thiếu chủ ngữ hoặc động từ; viết nhận định chắc chắn hơn điều nguồn cho phép. Cách sửa: chỉ rõ người nói, người nghe, nghĩa đang dùng và nguồn cho mỗi ý cần kiểm tra.'),PageBreak(),head('5 Reading and Comprehension'),para(l['readingTitle'],'v'),bilingual(l['readingPairs']),para('Source note: Original educational passage for this unit, informed by '+l['reference']+'. The Vietnamese column explains this authored passage; neither column is an official ecclesial quotation or prayer translation.','s'),head('Reading guide'),bilingual([('Paragraphs 1–2 introduce the concepts and their limits. Paragraph 3 gives a particular situation. A detail from the situation illustrates a point; it does not prove that all communities have the same arrangements.','Đoạn 1–2 nêu khái niệm và giới hạn. Đoạn 3 đưa tình huống cụ thể. Chi tiết minh họa một ý, không chứng minh mọi cộng đoàn đều có cùng cách sắp xếp.')]),para('Complete R01–R05 and C01 in the question paper or Word workbook. During final recall, try explaining the main point before reopening the reading.','s'),PageBreak(),head('6 Speaking and Writing Practice'),table([['Task code','Task','Output'],['S01',l['focus'],'100–130 words; prepare 2 minutes; rehearse 60–90 seconds.'],['S02','Learner or volunteer and English-speaking visitor','6 turns; clarify a question, explain a term, and identify a next step.'],['W01','Purposeful explanation or communication in the assigned situation','150–180 words; accurate reference; appropriate reader and purpose.']],[W*.13,W*.48,W*.39]),bilingual([('Use this structure: purpose or main point, explanation, a concrete example, a source where relevant, and a suitable conclusion or next step. Make the Catholic distinction accurate before attempting a more complex sentence.','Có thể dùng bố cục: mục đích hoặc ý chính, giải thích, ví dụ cụ thể, nguồn khi cần, kết luận hoặc bước tiếp theo. Làm đúng phân biệt Công giáo trước khi cố dùng câu phức tạp.')]),para('Support: '+l['support']),para('Extension: '+l['stretch']),head('Listening practice'),para('A downloadable recording accompanies L01–L03. Listen twice before viewing the transcript in the appendix. Audio uses the Flite SLT synthetic American English voice, not a human speaker.'),para('https://catholic-english-study.calmgnome10.chatgpt.site/'+l['audio'],'s'),head('7 Extended Catholic and Jesuit Knowledge')]
 for x in l['extension']:s +=[para(x['theme'],'v'),bilingual([(x['en'],x['vi'])]),para(x['status'],'s')]
 s +=[head('8 Summary and Active Recall')]+[para(f'{i:02}. {x}') for i,x in enumerate(l['recall'],1)]
 s +=[bilingual([('Close the textbook and explain three terms, two distinctions, and one practical application. Then check your explanation against the lesson, correct two errors, and repeat after one, three, and seven days.','Đóng giáo trình, giải thích ba từ, hai phân biệt và một cách áp dụng. Sau đó đối chiếu, sửa hai lỗi rồi ôn sau một, ba và bảy ngày.')]),head('References and Academic Integrity'),para('Syllabus Catholic Vocabulary and Ecclesial Language 2026–2027, pp. 1–4. Content in this unit is supplementary, not a prescribed reading supplied by the instructor.','s'),para(l['reference'],'s')]
 for title,url in source_items(l):s +=[para(title,'s'),para(url,'s')]
 s +=[bilingual([('Follow the syllabus AI-limited policy. Do not use this material to replace a restricted personal reflection, in-class vocabulary test, or another assignment designated AI-limited. Original writing and speaking practice should begin with your own attempt.','Tuân thủ chính sách AI-limited. Không dùng học liệu thay bài suy tư cá nhân, kiểm tra từ vựng trên lớp hoặc bài bị hạn chế AI. Bắt đầu viết và nói bằng phần tự làm của bạn.')]),PageBreak(),head('Appendix Answer Guidance and Rubrics'),para('Chỉ xem sau khi hoàn thành hoặc chủ động chọn xem. Đáp án tự luận là các ý cần có, không phải một đoạn văn duy nhất bắt buộc học thuộc. Điểm dưới đây là tự học, không phải trọng số thi chính thức.','s')]
 for q in l['questions']:
  ans=chr(65+q['answer'])+'. '+q['options'][q['answer']] if q['type']=='mc' else q['answer']
  s +=[head(q['code']+' '+str(q['points'])+' points'),para(ans),para(q['explanation'],'s')]
  for r in q['criteria']:s +=[para(r['label']+' /'+str(r['max'])+': '+r['desc'],'s')]
 s +=[head('Listening Transcript'),para(l['script']),para('Audio is synthetic Flite SLT speech. It is an original learning scenario, not a reproduction of an official prayer or liturgical text.','s'),head('Separate Content and Language Feedback'),para('Ghi riêng điểm nội dung/thuật ngữ Công giáo, điểm ngôn ngữ và điểm hiểu thông tin nghe theo rubric từng câu. Bản viết cho phần Speaking chỉ đánh giá được kịch bản; cần nghe thực tế để nhận xét phát âm. Máy chỉ chấm trắc nghiệm; tự luận cần người học hoặc giáo viên chấm.','s')]
 writepdf(D/l['downloads'][0],s,str(n))

def questions(l):
 n=l['number'];s=[para(f'UNIT {n:02} ASSESSMENT WORKSHEET','t'),para('Catholic English B1 core and optional extension','s'),para(f'CE-L{n:02} • Đề bài dành cho in ấn • 19 nhiệm vụ • 100 điểm đề xuất','s'),para('Họ tên / Full name: __________________________  Ngày / Date: __________'),table(marking(l),[W*.62,W*.23,W*.15]),para('Answer independently in English. Reading and audio resources are provided here. The matching Word workbook has the same codes, wording, choices, and reading, with ruled space for answers. No answers are included.','s')]
 for a,skill,name,codes in GROUPS:
  s +=[PageBreak() if a=='A' else Spacer(1,6),head('TASK '+a+' '+name.upper()+' '+str(sum(q['points'] for q in l['questions'] if q['skill']==skill))+' POINTS')]
  if skill=='Reading':s +=[para(l['readingTitle'],'v'),para(l['reading']),para('Choose one best answer A, B, C, or D.','s')]
  if skill=='Terminology':s +=[para('Answer each question in 40–70 English words. Explain meanings, distinguish the terms, and include a relevant example.')]
  if skill=='Speaking':s +=[para('Use the preparation and speaking times in each task. A written script alone does not establish pronunciation ability.')]
  if skill=='Writing':s +=[para('Use the stated reader, purpose, and length. Do not write a personal testimony or invent an official quotation.')]
  if skill=='Listening':s +=[para('Audio: Flite SLT synthetic American English voice. Listen twice before consulting the transcript, which appears only in the curriculum appendix.'),para('https://catholic-english-study.calmgnome10.chatgpt.site/'+l['audio'],'s')]
  s +=[qpdf(q) for q in l['questions'] if q['skill']==skill]
 s +=[head('Assessment Criteria'),para('Catholic content and terminology are distinguished from language, task completion, and source use. See the curriculum appendix for the criterion allocations. A correct paraphrase need not match a model word for word.'),para('This is an independent self-study assessment; it is not the official course examination.','s')]
 writepdf(D/l['downloads'][1],s,str(n))

# Word follows the supplied worksheet's typography and ruled writing space, on A4.
def doc_table(doc,rows,widths,header=True):
 t=doc.add_table(rows=0,cols=len(widths));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 for j,w in enumerate(widths):t.columns[j].width=Inches(w)
 for i,row in enumerate(rows):
  cells=t.add_row().cells
  for j,text in enumerate(row):
   cells[j].width=Inches(widths[j]);cells[j].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   p=cells[j].paragraphs[0];r=p.add_run(str(text));r.font.size=Pt(11);r.bold=bool(header and i==0);p.paragraph_format.space_after=Pt(4);p.paragraph_format.space_before=Pt(4)
   pr=cells[j]._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');b.append(e)
   pr.append(b);mar=OxmlElement('w:tcMar')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:w'),'90');e.set(qn('w:type'),'dxa');mar.append(e)
   pr.append(mar)
   if header and i==0:
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EEEEEE');pr.append(sh)
  trpr=t.rows[-1]._tr.get_or_add_trPr();ns=OxmlElement('w:cantSplit');trpr.append(ns)
  if header and i==0:trpr.append(OxmlElement('w:tblHeader'))
 doc.add_paragraph().paragraph_format.space_after=Pt(3)
 return t

def lines(doc,num):
 for i in range(num):
  p=doc.add_paragraph();p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=Pt(20)
  p.add_run('_'*70)

def practice(l):
 n=l['number'];d=Document();sec=d.sections[0];sec.page_width=Inches(8.2677);sec.page_height=Inches(11.6929);sec.top_margin=Inches(.68);sec.bottom_margin=Inches(.65);sec.left_margin=Inches(.72);sec.right_margin=Inches(.72)
 for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
  st=d.styles[name];st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0);st.font.size=Pt(13 if name=='Normal' else 20 if name=='Title' else 16 if name.startswith('Heading') else 11);st.paragraph_format.space_after=Pt(7)
  st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Times New Roman')
 d.styles['Normal'].paragraph_format.line_spacing=1.1
 d.add_paragraph(f'UNIT {n:02} ASSESSMENT WORKSHEET','Title');d.add_paragraph(f'Catholic English B1 core and optional extension | CE-L{n:02} | Phiếu có dòng viết đáp án','Subtitle')
 d.add_paragraph('Họ tên / Full name: __________________________  Ngày / Date: __________');d.add_paragraph('Bắt đầu / Start: __________________  Kết thúc / Finish: __________________')
 d.add_paragraph('Bài tự học gồm 19 nhiệm vụ, 100 điểm đề xuất. Không kèm đáp án. Mã câu và nội dung khớp đề PDF; chỉ phần chỗ viết được bổ sung. Không phải đề thi chính thức.')
 doc_table(d,marking(l),[4.15,1.6,1.03])
 for a,skill,name,codes in GROUPS:
  d.add_page_break();d.add_paragraph('TASK '+a+' '+name.upper(), 'Heading 1')
  if skill=='Reading':
   d.add_paragraph(l['readingTitle'],'Heading 2')
   for p in l['reading'].split('\n\n'):d.add_paragraph(p)
   d.add_paragraph('Choose ONE best answer A, B, C, or D. A and B share the first row; C and D share the second.')
  if skill=='Terminology':d.add_paragraph('Write 40–70 English words for each question. Explain meanings, distinguish the terms, and give a relevant example.')
  if skill=='Speaking':d.add_paragraph('Write your script or dialogue, then practise aloud. Pronunciation cannot be graded from written responses alone. Use the preparation and performance times in the tasks.')
  if skill=='Writing':d.add_paragraph('Write for the stated reader and purpose. Use your own draft and an accurate reference; do not invent an official quotation.')
  if skill=='Listening':
   d.add_paragraph('Listen twice before reading the transcript in the curriculum appendix. Synthetic American English voice, Flite SLT.');d.add_paragraph('https://catholic-english-study.calmgnome10.chatgpt.site/'+l['audio'])
  for q in [q for q in l['questions'] if q['skill']==skill]:
   if q['code'] in ['V03','S02']:
    d.add_page_break();d.add_paragraph('TASK '+a+' continued','Heading 2')
   p=d.add_paragraph();p.paragraph_format.keep_with_next=True;p.add_run(f"{q['code']} | {q['prompt']} ({q['points']} points)").bold=True
   if q['type']=='mc':
    t=doc_table(d,[[chr(65+k)+'. '+q['options'][k] for k in [0,1]],[chr(65+k)+'. '+q['options'][k] for k in [2,3]]],[3.39,3.39],False)
    for ri,row in enumerate(t.rows):
     for ci,cell in enumerate(row.cells):
      for p in cell.paragraphs:p.paragraph_format.keep_with_next=not(ri==1 and ci==1)
    lines(d,1)
   else:
    count=8 if skill=='Terminology' else 3 if skill=='Grammar' else 14 if q['code']=='S01' else 12 if q['code']=='S02' else 24 if skill=='Writing' else 16 if skill=='Critical reading' else 5
    if skill=='Speaking':d.add_paragraph('Dàn ý / Key points and follow-up:');lines(d,3)
    lines(d,count)
    if skill=='Writing':d.add_paragraph('Số từ / Word count: __________')
 d.add_paragraph('ASSESSMENT CRITERIA','Heading 1');d.add_paragraph('Nội dung Công giáo, dùng nguồn và năng lực ngôn ngữ được phân biệt theo rubric ở phụ lục giáo trình. Không có đáp án trong phiếu này. Đây là bài tự học, không thay bài AI-limited của lớp.')
 foot=sec.footer.paragraphs[0];foot.add_run(f'Catholic English | Unit {n:02} | CE-L{n:02} | ')
 fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');foot._p.append(fld)
 for r in foot.runs:r.font.size=Pt(9)
 d.save(D/l['downloads'][2])

if __name__=='__main__':
 import sys
 select=int(sys.argv[1]) if len(sys.argv)>1 else None
 for l in C['lessons'][1:]:
  if select is not None and l['number']!=select:continue
  curriculum(l);questions(l);practice(l);print('Created Unit',l['number'],flush=True)
